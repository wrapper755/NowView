"""
SC01 규칙 계산 — 엑셀의 **파생 시트를 기대값으로** 대조한다(DATA-01·TEST-02).

시드 적재기는 원천 시트만 읽고, 이 테스트는 파생 시트(월분기연간·SC01검증)를 따로 읽어
서비스가 다시 계산한 값과 맞춘다. 숫자를 테스트에 손으로 옮겨 적지 않는다 — 엑셀이 고쳐지면
테스트도 따라오고, 계산이 틀리면 바로 드러난다. 예외 사례(EX)·경계값(AT06)은 엑셀에 계산이
없으므로 구현명세의 기대값을 그대로 적었다.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.sc01.calc import (
    PERIODS,
    Risk,
    assess,
    assess_all,
    build_forecast,
    monthly_pnl,
    period_pnl,
)
from app.sc01.seed import SEED_PATH, Key, load_seed

SEED = load_seed()
CUS, ULS = "CUS-HK", "ULS"
LABEL = {"기준 이내": Risk.NORMAL, "관심": Risk.ATTENTION, "중요": Risk.CRITICAL}


def k(model: str, month: int) -> Key:
    return (CUS, ULS, f"A-{model}", date(2027, month, 1))


def derived(sheet: str, first: str) -> list[dict]:
    ws = load_workbook(SEED_PATH, data_only=True, read_only=True)[sheet]
    header, out = None, []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = list(row) if row and row[0] == first else None
            continue
        if any(v is not None for v in row):
            out.append(dict(zip(header, row, strict=False)))
    return out


def f02() -> dict[Key, int | None]:
    return build_forecast(SEED.plans["F01"], SEED.customer_plans["CP02"])


# ─────────────────────────── 전망 생성 ───────────────────────────


def test_generated_f02_equals_the_seeded_f02() -> None:
    """V10·V11 — F02는 F01 복제 + CP02의 7~9월. 1~6월·10~12월은 F01 그대로다."""
    assert f02() == SEED.plans["F02"]


def test_only_july_b_changes() -> None:
    diff = {key for key, q in f02().items() if q != SEED.plans["F01"][key]}
    assert diff == {k("b", 7)}


def test_the_40_percent_share_is_never_applied() -> None:
    """원천이 이미 당사 배정물량이다(G10). 고객 전체 물량이 오면 변환 규칙 없이 쓰지 않는다."""
    row = SEED.customer_plans["CP02"][0]
    from dataclasses import replace

    with pytest.raises(ValueError, match="변환 규칙"):
        build_forecast(SEED.plans["F01"], [replace(row, allocation_scope="CUSTOMER_TOTAL")])


# ─────────────────────────── 지표·위험 (SC01검증 시트) ───────────────────────────


@pytest.mark.parametrize(
    "row", derived("SC01검증", "run_id"), ids=lambda r: f"{r['model_id']}-{r['period']:%m}"
)
def test_metrics_match_the_validation_sheet(row: dict) -> None:
    """AT03·AT04 — 15행 전부. 비율은 8자리 반올림값으로 대조한다."""
    key = (CUS, ULS, row["model_id"], row["period"].date())
    a = assess_all([key], SEED.plans["T01"], SEED.plans["F01"], f02(), SEED)[0]
    m = a.metrics
    q = Decimal("0.00000001")

    assert m["KPI01"].value == Decimal(str(row["달성률"])).quantize(q)
    assert m["M01"].value == row["추가감소"]
    assert m["KRI01"].value == Decimal(str(row["감소율"])).quantize(q)
    assert m["M03"].value == row["순손익영향"]
    assert m["KRI02"].value == row["감소액"]
    assert m["KRI03"].value == Decimal(str(row["목표부족률"])).quantize(q)
    assert a.risk is LABEL[row["위험수준"]]


def test_the_base_case() -> None:
    """AT03 — b 7월: 700 / 7.368421% / 14,000,000원 / 12% / 중요."""
    a = assess_all([k("b", 7)], SEED.plans["T01"], SEED.plans["F01"], f02(), SEED)[0]
    assert a.metrics["M01"].value == 700
    assert a.metrics["KRI01"].value == Decimal("0.07368421")
    assert a.metrics["KRI02"].value == 14_000_000
    assert a.metrics["KRI03"].value == Decimal("0.12")
    assert a.metrics["M02"].value == 1200  # 목표 부족물량 ≠ 이번 추가 감소
    assert a.risk is Risk.CRITICAL
    # 세 지표가 모두 중요다 — Agent 기대 설명(구현명세 11장)이 이것을 인용한다.
    assert {a.metrics[x].risk for x in ("KRI01", "KRI02", "KRI03")} == {Risk.CRITICAL}


def test_a_stays_attention_from_the_existing_shortfall() -> None:
    """D13 — a는 변경이 없어도 목표 부족률 5%로 관심이다(기존 이슈 유지)."""
    a = assess_all([k("a", 7)], SEED.plans["T01"], SEED.plans["F01"], f02(), SEED)[0]
    assert a.metrics["M01"].value == 0
    assert a.risk is Risk.ATTENTION


# ─────────────────────────── 경계값 (AT06) ───────────────────────────

R = SEED.rules


@pytest.mark.parametrize(
    ("prior", "latest", "cm", "metric", "expected"),
    [
        # KRI01 감소율 3%·5%: 경계 이상은 그 수준, 직하는 아래 수준
        (10000, 9700, 0, "KRI01", Risk.ATTENTION),
        (10000, 9701, 0, "KRI01", Risk.NORMAL),
        (10000, 9500, 0, "KRI01", Risk.CRITICAL),
        (10000, 9501, 0, "KRI01", Risk.ATTENTION),
        # KRI02 감소액 5백만·1천만원
        (10000, 9750, 20000, "KRI02", Risk.ATTENTION),
        (10000, 9751, 20000, "KRI02", Risk.NORMAL),
        (10000, 9500, 20000, "KRI02", Risk.CRITICAL),
        (10000, 9501, 20000, "KRI02", Risk.ATTENTION),
    ],
)
def test_threshold_boundaries(
    prior: int, latest: int, cm: int, metric: str, expected: Risk
) -> None:
    a = assess(k("b", 7), None, prior, latest, Decimal(cm), R)
    assert a.metrics[metric].risk is expected


@pytest.mark.parametrize(
    ("target", "latest", "expected"),
    [
        (100, 95, Risk.ATTENTION),
        (100, 96, Risk.NORMAL),
        (100, 90, Risk.CRITICAL),
        (100, 91, Risk.ATTENTION),
    ],
)
def test_shortfall_boundaries(target: int, latest: int, expected: Risk) -> None:
    a = assess(k("b", 7), target, latest, latest, Decimal(0), R)
    assert a.metrics["KRI03"].risk is expected


# ─────────────────────────── 예외 (EX01~06) ───────────────────────────


def _one(prior: int | None, latest: int | None, model: str = "b", target: int | None = None):
    key = k(model, 7)
    t = SEED.plans["T01"][key] if target is None else target
    return assess(key, t, prior, latest, SEED.costs[key].cm, R)


def test_ex01_small_change_stays_within_limits() -> None:
    """c 5,940→5,900: 감소율 0.6734%, 감소액 112만원, 기준 이내."""
    a = _one(5940, 5900, "c")
    assert a.metrics["KRI01"].value == Decimal("0.00673401")
    assert a.metrics["KRI02"].value == 1_120_000
    assert a.metrics["KRI03"].value == Decimal("0.01666667")
    assert a.risk is Risk.NORMAL


def test_ex02_missing_quantity_holds_the_judgement() -> None:
    """누락은 0이 아니다 — 판정 보류, 지표 미산출."""
    a = _one(9500, None)
    assert a.risk is None
    assert not a.metrics["KRI01"].computable and a.metrics["KRI01"].reason == "입력 누락"


def test_ex03_correction_is_judged_against_the_approved_prior() -> None:
    """정정 9,000: 공식 비교는 F01(9,500) — 500 감소, 1천만원 → 중요."""
    a = _one(9500, 9000)
    assert a.metrics["M01"].value == 500
    assert a.metrics["KRI02"].value == 10_000_000
    assert a.risk is Risk.CRITICAL


def test_ex04_offset_keeps_the_model_alarm() -> None:
    """c +1,568만원이 b −1,400만원을 상쇄해도 b는 중요다(G14)."""
    b, c = _one(9500, 8800), _one(5940, 6500, "c")
    assert c.metrics["M03"].value == 15_680_000
    assert b.metrics["M03"].value + c.metrics["M03"].value == 1_680_000
    assert b.risk is Risk.CRITICAL
    assert c.metrics["KRI02"].value == 0 and c.metrics["KRI01"].value < 0  # 증가 부호 보존


def test_ex05_confirmed_zero_is_a_full_decrease() -> None:
    a = _one(9500, 0)
    assert a.metrics["KRI01"].value == 1
    assert a.metrics["KRI02"].value == 190_000_000
    assert a.risk is Risk.CRITICAL


def test_ex06_zero_prior_is_not_computable_but_others_are() -> None:
    a = _one(0, 100, target=100)
    assert not a.metrics["KRI01"].computable
    assert a.metrics["KRI01"].reason == "직전 전망 0 — 미산출"
    assert a.metrics["M03"].value > 0  # 절대 변화·손익은 표시한다
    assert a.partial and a.risk is Risk.NORMAL


# ─────────────────────────── 손익 집계 (월분기연간 시트) ───────────────────────────

AGG = derived("월분기연간", "기간")


@pytest.mark.parametrize("row", AGG, ids=lambda r: r["기간"])
def test_pnl_matches_the_period_sheet(row: dict) -> None:
    """AT05·V04~V08·V12 — 월·분기·연간 매출·공헌이익·관리손익, 사건 순영향, 목표 차이."""
    start, end = PERIODS[row["기간"]]
    t, p1, p2 = (
        period_pnl(monthly_pnl(plan, SEED), start, end)
        for plan in (SEED.plans["T01"], SEED.plans["F01"], f02())
    )
    assert (t.revenue, p1.revenue, p2.revenue) == (row["T매출"], row["F01매출"], row["F02매출"])
    assert (t.cm, p1.cm, p2.cm) == (row["T공헌이익"], row["F01공헌이익"], row["F02공헌이익"])
    assert t.fixed == row["공장고정비"]
    assert (t.mgmt_profit, p1.mgmt_profit, p2.mgmt_profit) == (
        row["T관리손익"],
        row["F01관리손익"],
        row["F02관리손익"],
    )
    assert p2.mgmt_profit - p1.mgmt_profit == row["사건순영향"]
    assert t.mgmt_profit - p2.mgmt_profit == row["목표차이"]


def test_annual_shortfall_is_separate_from_the_event() -> None:
    """AT07 — 연간 목표 부족액 167,939,000원 ≠ 이번 사건 감소액 14,000,000원."""
    months = {n: monthly_pnl(SEED.plans[n], SEED) for n in ("T01", "F01")}
    months["F02"] = monthly_pnl(f02(), SEED)
    year = {n: period_pnl(m, 1, 12).mgmt_profit for n, m in months.items()}
    assert year["T01"] - year["F02"] == 167_939_000
    assert year["F01"] - year["F02"] == 14_000_000


def test_future_actuals_do_not_exist() -> None:
    """TEST-02 — 기본 시드에 6/15 이후 확정 실적이 없다. 6월은 누계다."""
    assert max(key[3] for key in SEED.actuals) == date(2027, 6, 1)
    assert all(c <= date(2027, 6, 15) for c in SEED.actual_cutoffs.values())
    assert len(SEED.actuals) == 30


def test_cost_identity() -> None:
    """V13 — 가격 − 변동비 = 기준 공헌이익."""
    assert all(c.price - c.variable_cost == c.cm for c in SEED.costs.values())
