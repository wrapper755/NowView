"""
SC05 규칙 계산 — 통합가상데이터 v0.5의 장면 5 파생 시트와 대조한다(값을 시험에 옮겨 적지 않는다).

* 기초 손익(과제 전) = SC05모델손익 60행 · 월 효과 = SC05월효과 18행 · 월 손익 = SC05월손익 12행 +
  연간·마감·전망·Q3
* 초기 지표 = SC05초기지표 · 예산 = SC05예산 · 변경 대사 = SC05변경대사 · 기존 45개 시트 = v0.4
* 경계 — 공유·확인 임계값(같으면 정상), 미확인은 평균에서 빼고 따로, 분모 0은 비해당, 중복 주문·비용
  미확인은 확정 차단
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.sc02.calc import Option
from app.sc04 import calc as sc04_calc
from app.sc04.seed import SEED_PATH as SC04_PATH
from app.sc04.seed import load_seed as load_sc04
from app.sc05 import calc
from app.sc05.seed import SEED_PATH, kst, load_seed

SEED = load_seed()
OPTION = Option(
    "C02-APP01", "APPROVED", "P-COM01", "SUP-01", Decimal(2000), Decimal(2120), date(2027, 7, 1)
)
FIXED = {m.isoformat(): str(SEED.numbers.fixed[m]) for m in calc.MONTHS}


def _assumptions() -> dict[str, dict]:
    """장면 4 판단본(DEC41-v2)의 가정 — ALT42는 질의 뒤 판단 시작(9월)."""
    rows = load_sc04().by_id("SC04가정")
    out = {task: sc04_calc.assumption_from_row(rows[alt]) for task, alt in calc.ALT_OF.items()}
    out["ACT42"]["start"] = rows["ALT42"]["판단 시작"].date().isoformat()
    return out


def _orders() -> list[dict]:
    return [
        {
            "order_id": o["order_id"],
            "qty": o["수량"],
            "urgent": bool(o["긴급여부"]),
            "extra": str(o["추가비합계"]),
            "late": o["납기이상"],
            "quality": o["품질이상"],
            "result_status": o["확인상태"],
        }
        for o in SEED.rows("SC05구매표본")
    ]


def _changes() -> list[dict]:
    return [
        {
            "change_id": c["change_id"],
            "received_at": kst(c["접수_KST"]),
            "shared_at": kst(c["공유_KST"]),
            "ack_at": kst(c["구매확인_KST"]),
        }
        for c in SEED.rows("SC05변경공유")
    ]


def _observation(orders=None, changes=None, support=Decimal(200000)) -> dict:
    return calc.observation(
        orders if orders is not None else _orders(),
        changes if changes is not None else _changes(),
        base_count=12,
        unit_cost=Decimal(1200000),
        support=support,
        share_limit=int(SEED.setting("share_minutes")),
        ack_limit=int(SEED.setting("confirm_minutes")),
    )


def _effects(nov: int = 3) -> list[dict]:
    """RUN51 입력 — 마감월은 예산 원장 집행·관측, 잔여월은 운영계획(10월 4, 11·12월 3)."""
    obs = _observation()
    settle = [t for t in SEED.rows("SC05예산원장") if t["유형"] == "SETTLE"]
    counts = {"2027-10-01": 4, "2027-11-01": nov, "2027-12-01": 3}
    out = []
    for task, a in _assumptions().items():
        for m in calc.EFFECT_MONTHS:
            i = calc.plan_inputs(a, m)
            if m < date(2027, 10, 1):
                spent = [
                    t for t in settle if t["task_id"] == task and t["발생_KST"].month == m.month
                ]
                i["one_off"] = str(
                    sum(t["손익비용"] for t in spent if t["budget_scope"] == "ONEOFF")
                )
                i["support"] = str(
                    sum(t["손익비용"] for t in spent if t["budget_scope"] == "OPERATING")
                )
                if task == "ACT42" and i["applied"]:
                    i.update(count=obs["urgent"], extra_cost=obs["extra_cost"])
            elif task == "ACT42":
                i["count"] = counts[m.isoformat()]
            out.append(calc.effect(task, i))
    return out


def _prior() -> tuple[list[dict], list[dict]]:
    effects = [
        calc.effect(t, calc.plan_inputs(a, m))
        for t, a in _assumptions().items()
        for m in calc.EFFECT_MONTHS
    ]
    return calc.months(
        calc.basis(SEED.numbers, OPTION), effects, FIXED, first_forecast=date(2027, 9, 1)
    ), effects


def test_v05_keeps_the_45_original_sheets() -> None:
    """v0.5는 기존 시트를 고치지 않았다 — 숫자 원천은 장면 2~4와 같다."""
    old = openpyxl.load_workbook(SC04_PATH, data_only=True, read_only=True)
    new = openpyxl.load_workbook(SEED_PATH, data_only=True, read_only=True)
    for name in old.sheetnames:
        assert list(old[name].iter_rows(values_only=True)) == list(
            new[name].iter_rows(values_only=True)
        ), name
    assert sum(1 for n in new.sheetnames if n.startswith("SC05")) == 26


def test_basis_is_the_model_profit_sheet() -> None:
    """
    과제 전 공헌이익·Target 공헌이익·물량·매출 60행 — 7~9월 가상 마감 기초값은 기존 전망과 같다.
    """
    base = {(b["period"][:7], b["model_id"]): b for b in calc.basis(SEED.numbers, OPTION)}
    rows = SEED.rows("SC05모델손익")
    assert len(rows) == 60
    for r in rows:
        b = base[(r["월"].date().isoformat()[:7], r["model_id"])]
        assert (
            b["qty"],
            Decimal(b["basis_cm"]),
            Decimal(b["target_cm"]),
            Decimal(b["revenue"]),
        ) == (
            r["물량"],
            Decimal(r["과제전공헌이익"]),
            Decimal(r["Target공헌이익"]),
            Decimal(r["매출"]),
        ), r["row_id"]


def test_prior_official_fc50_is_the_sc04_decision() -> None:
    """FC50 월 손익 = 장면 4 판단 가정으로 낸 손익(DEC41-v2와 수치가 같다 — 별도 승인일 뿐이다)."""
    prior, _ = _prior()
    fc50 = {r["월"].date().isoformat(): Decimal(r["관리손익"]) for r in SEED.rows("SC05이전공식")}
    assert {m["period"]: Decimal(m["profit"]) for m in prior} == fc50
    assert sum(fc50.values()) == 3427786680


def test_monthly_effects_match_the_effect_sheet() -> None:
    """
    SC05월효과 18행 — ACT41 회피 변동가공비, ACT42 긴급비용 차이(9월은 관측), ACT43 대체단가. 고정비
    절감 0.
    """
    got = {(e["task_id"], e["month"][:7]): e for e in _effects()}
    for r in SEED.rows("SC05월효과"):
        e = got[(r["task_id"], r["월"].date().isoformat()[:7])]
        assert Decimal(e["net"]) == Decimal(r["순효과"]), r["effect_id"]
        assert int(e["qty"]) == r["적용수량"] and int(e["base_count"]) == r["기준긴급건"], r[
            "effect_id"
        ]
        assert int(e["count"]) == r["실제·전망긴급건"] and Decimal(e["support"]) == Decimal(
            r["운영지원비"]
        ), r["effect_id"]
        assert Decimal(e["one_off"]) == Decimal(r["일회성비"]) and r["고정비절감"] == 0, r[
            "effect_id"
        ]


def test_monthly_profit_and_periods_match_the_sheet() -> None:
    """SC05월손익 — 월 12행과 연간·마감 1~9월·전망 10~12월·Q3. 원 단위 오차 0(T520)."""
    rows = calc.months(
        calc.basis(SEED.numbers, OPTION), _effects(), FIXED, first_forecast=date(2027, 10, 1)
    )
    by = {m["period"]: m for m in rows}
    prior = {m["period"]: m for m in _prior()[0]}
    sheet = SEED.rows("SC05월손익")
    for r in sheet[:12]:
        m = by[r["기간"].date().isoformat()]
        assert m["kind"] == r["구분"]
        for key, col in (
            ("target", "Target"),
            ("basis_profit", "과제전관리손익"),
            ("effect", "과제순효과"),
            ("profit", "갱신관리손익"),
            ("gap", "목표부족"),
            ("revenue", "매출"),
            ("fixed", "공장고정비"),
        ):
            assert Decimal(m[key]) == Decimal(r[col]), (r["기간"], key)
        assert Decimal(prior[m["period"]]["profit"]) == Decimal(r["장면4판단전망"])
    p = calc.periods(rows)
    for kind, label in (
        ("year", "2027 연간"),
        ("closed", "2027 마감1~9월"),
        ("forecast", "2027 전망10~12월"),
        ("q3", "2027 Q3"),
    ):
        r = next(x for x in sheet if x["기간"] == label)
        assert (
            Decimal(p[kind]["profit"]),
            Decimal(p[kind]["effect"]),
            Decimal(p[kind]["gap"]),
        ) == (
            Decimal(r["갱신관리손익"]),
            Decimal(r["과제순효과"]),
            Decimal(r["목표부족"]),
        ), kind
    assert (p["year"]["profit"], p["year"]["gap"], p["year"]["fixed"]) == (
        "3428986680",
        "48303320",
        "4440000000",
    )


def test_model_effects_are_added_once_to_the_linked_model() -> None:
    """
    모델별 과제 순효과 — ACT41은 A-a, ACT42·43은 A-b. c~e는 0. 고정비는 모델에 배부하지 않는다.
    """
    rows = calc.months(
        calc.basis(SEED.numbers, OPTION), _effects(), FIXED, first_forecast=date(2027, 10, 1)
    )
    got = {(m["period"][:7], r["model_id"]): r for m in rows for r in m["models"]}
    for r in SEED.rows("SC05모델손익"):
        g = got[(r["월"].date().isoformat()[:7], r["model_id"])]
        assert (Decimal(g["effect"]), Decimal(g["cm"])) == (
            Decimal(r["과제순효과"]),
            Decimal(r["갱신공헌이익"]),
        ), r["row_id"]


def test_reconciliation_against_fc50_explains_every_won() -> None:
    """
    SC05변경대사 — 9월 −1.2M(실제 5건), 10월 0, 11~12월 +2.4M, 그 외 0, 대사차이 0(T521의 반대쪽).
    """
    rows = calc.months(
        calc.basis(SEED.numbers, OPTION), _effects(), FIXED, first_forecast=date(2027, 10, 1)
    )
    prior, prior_eff = _prior()
    r = calc.reconcile(rows, _effects(), prior, prior_eff)
    sheet = {x["원인"]: Decimal(x["전망변화"]) for x in SEED.rows("SC05변경대사")}
    assert (
        Decimal(r["prior_total"]) == sheet["종전 공식 FC50"]
        and Decimal(r["new_total"]) == sheet["변경후 전망"]
    )
    by_month = {x["month"][:7]: Decimal(x["amount"]) for x in r["rows"] if x["cause"] == "urgent"}
    assert by_month.get("2027-09") == sheet["9월 긴급구매 실제 5건 vs 전망4건"]
    assert by_month.get("2027-10", Decimal(0)) == sheet["10월 긴급구매 전망4건 유지"]
    assert by_month["2027-11"] + by_month["2027-12"] == sheet["11~12월 긴급구매 전망3건"]
    assert [x for x in r["rows"] if x["cause"] != "urgent"] == []
    assert Decimal(r["residual"]) == sheet["RUN51 대사차이"] == 0 and r["ok"]


def test_x506_re_adding_the_approved_effect_fails_reconciliation() -> None:
    """
    X506·T521 — 종전 공식본에 장면 4 승인효과 127.8M을 다시 더하면 원인으로 설명되지 않는 차이가
    남는다.
    """
    rows = calc.months(
        calc.basis(SEED.numbers, OPTION), _effects(), FIXED, first_forecast=date(2027, 10, 1)
    )
    prior, prior_eff = _prior()
    r = calc.reconcile(
        rows, _effects(), prior, prior_eff, reported_total=Decimal(3427786680) + Decimal(127800000)
    )
    assert not r["ok"] and Decimal(r["residual"]) == Decimal(126600000)


def test_x508_november_four_urgent_lowers_the_forecast() -> None:
    """X508·T525 — 11월 3→4건: 연간 3,427,786,680원·부족 49,503,320원. 마감 실적은 그대로다."""
    rows = calc.months(
        calc.basis(SEED.numbers, OPTION), _effects(nov=4), FIXED, first_forecast=date(2027, 10, 1)
    )
    p = calc.periods(rows)
    assert (p["year"]["profit"], p["year"]["gap"]) == ("3427786680", "49503320")
    assert p["closed"]["profit"] == "2521317800"


def test_initial_metrics_match_the_sheet() -> None:
    """
    SC05초기지표 — 60건·9,785개·긴급 5건(8.333…%)·6백만원·순개선 8.2백만원·납기·품질 0·공유
    26.67분·확인 62.5분·지연 1.
    """
    obs = _observation()
    m = {r["지표ID"]: r["값"] for r in SEED.rows("SC05초기지표")}
    assert (obs["orders"], obs["urgent"], obs["qty"]) == (m["M505a"], m["M505b"], m["M505q"])
    assert Decimal(obs["urgent_rate"]) == Decimal(5) / Decimal(60)
    assert abs(float(obs["urgent_rate"]) - m["M505c"]) < 1e-12
    assert Decimal(obs["extra_cost"]) == Decimal(m["M506"]) and Decimal(
        obs["baseline_cost"]
    ) == Decimal(m["M507b"])
    assert (
        Decimal(obs["net"])
        == Decimal(m["M507"])
        == Decimal(14400000) - Decimal(6000000) - Decimal(200000)
    )
    assert (obs["late"], obs["quality"]) == (m["M508a"], m["M508b"])
    assert abs(float(obs["share"]["average"]) - m["M503"]) < 1e-9 and Decimal(
        obs["ack"]["average"]
    ) == Decimal(str(m["M504"]))
    assert obs["ack"]["delayed"] == m["M504d"] and obs["changes"] == 6 and obs["complete"]


def test_threshold_equal_is_on_time_and_null_is_unconfirmed() -> None:
    """
    X503·T514 — 정확히 120분은 정상, 넘으면 지연. 확인 NULL은 평균에서 빼고 미확인으로 센다(0분
    아님).
    """
    base = {"received_at": "2027-09-17T09:00:00+09:00", "shared_at": "2027-09-17T09:30:00+09:00"}
    exact = {**base, "change_id": "A", "ack_at": "2027-09-17T11:30:00+09:00"}
    over = {**base, "change_id": "B", "ack_at": "2027-09-17T11:30:01+09:00"}
    none = {**base, "change_id": "C", "ack_at": None}
    obs = _observation(changes=[exact, over, none])
    assert (
        obs["ack"]["delayed"] == 1
        and obs["ack"]["unconfirmed"] == 1
        and obs["ack"]["measured"] == 2
    )
    assert Decimal(obs["ack"]["average"]) == (Decimal(120) + Decimal(7201) / 60) / 2


def test_x509_duplicates_and_unknown_cost_block_confirmation() -> None:
    """
    X509·T516 — 같은 주문 ID 두 번, 긴급 추가비 NULL. 0원 정상으로 처리하지 않고 확정을 막는다.
    """
    orders = _orders()
    orders.append(dict(orders[1]))
    orders[9] = {**orders[9], "extra": None}  # PO51-010 긴급
    obs = _observation(orders=orders)
    assert obs["duplicates"] == ["PO51-002"] and obs["unknown_cost"] == ["PO51-010"]
    assert obs["extra_cost"] is None and obs["net"] is None and not obs["complete"]
    assert obs["orders"] == 60  # 중복은 한 번만 센다


def test_unconfirmed_results_are_not_zero_and_empty_scope_is_not_applicable() -> None:
    """T518 — 납기·품질 미수집이면 정상 0건이 아니라 미확인. 분모 0은 0%가 아니라 비해당."""
    orders = [{**o, "result_status": None} for o in _orders()]
    obs = _observation(orders=orders)
    assert obs["late"] is None and obs["quality"] is None and obs["unconfirmed_results"] == 60
    empty = _observation(orders=[], changes=[])
    assert (
        empty["urgent_rate"] is None and empty["ack"]["average"] is None and not empty["complete"]
    )


def test_budget_settle_does_not_double_count() -> None:
    """
    SC05예산 · T508 — COMMIT은 약정, SETTLE은 같은 약정을 줄이며 집행. 일회성 11백만원, 운영지원은
    별도 원장.
    """
    ledger = [
        {
            "task_id": t["task_id"],
            "scope": t["budget_scope"],
            "commitment_delta": str(t["약정증감"]),
            "paid_delta": str(t["집행증감"]),
            "pnl_cost": str(t["손익비용"]),
        }
        for t in SEED.rows("SC05예산원장")
    ]
    limits = {"ACT41": Decimal(8000000), "ACT42": Decimal(6000000), "ACT43": Decimal(4000000)}
    b = calc.budget(limits, Decimal(800000), ledger)
    sheet = SEED.by_id("SC05예산")
    for t in b["tasks"]:
        r = sheet[t["task_id"]]
        assert (
            Decimal(t["limit"]),
            Decimal(t["commitment"]),
            Decimal(t["paid"]),
            Decimal(t["available"]),
        ) == (
            Decimal(r["일회성 승인한도"]),
            Decimal(r["미집행약정"]),
            Decimal(r["집행누계"]),
            Decimal(r["사용가능액"]),
        )
    assert (b["total"]["paid"], b["total"]["available"], b["total"]["commitment"]) == (
        "7000000",
        "11000000",
        "0",
    )
    assert (b["operating"]["commitment"], b["operating"]["paid"]) == ("600000", "200000")


@pytest.mark.parametrize("plan_ref", ["CP51", "CP52"])
def test_customer_plans_keep_the_month_total(plan_ref: str) -> None:
    """T510 — a~e 전부, 각 월합계 변화 0. 두 요청창은 각각 비교안이다(합산하지 않는다)."""
    from app.sc05.domain import customer_plan

    lines = customer_plan(SEED, plan_ref)
    assert [ln["model_id"] for ln in lines] == ["A-a", "A-b", "A-c", "A-d", "A-e"]
    assert all(ln["month_diff"] == 0 for ln in lines)
    assert [ln["model_id"] for ln in lines if ln["changed"]] == ["A-b"]
    assert sum(ln["month_total"] for ln in lines if ln["model_id"] == "A-b") == 9785
