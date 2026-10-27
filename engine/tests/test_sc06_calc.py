"""
SC06 계산 — 통합가상데이터 v0.6의 장면 6 파생
시트(과제성과·월효과·원인집계·운영지표·분기손익·검증)와 대조한다. 값을
시험에 옮겨 적지 않고 시트에서 읽는다. 기대(BASE61)는 장면 5를 실제로 끝낸 FC51 산정본과
대사한다(CK61-14).

예외는 구현명세 12장 AT61-06~12 · SC06예외 X61-01~08이다 — 시드의 복제본에 변경을 넣어
적재한다(기본 적재에 섞지 않는다).
"""

from __future__ import annotations

import copy
import sqlite3
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.sc01.domain import DomainError
from app.sc01.process import Sc01Process
from app.sc01.store import SqliteSc01Store
from app.sc02.process import Sc02Process
from app.sc03.process import Sc03Process
from app.sc04.process import Sc04Process
from app.sc05.process import Sc05Process
from app.sc05.seed import SEED_PATH as SC05_PATH
from app.sc06 import calc, domain
from app.sc06.seed import SEED_PATH, kst, load_seed

SEED = load_seed()
PREFIXES = ("sc01", "sc02", "sc03", "sc04", "sc05")


@pytest.fixture(scope="module")
def sc05_doc() -> dict:
    """장면 5 기본 시연을 끝까지(HANDOFF51) — FC51 산정본이 장면 6 기대의 출처다."""
    stores = [
        SqliteSc01Store(sqlite3.connect(":memory:", check_same_thread=False), p) for p in PREFIXES
    ]
    sc02 = Sc02Process(stores[1], Sc01Process(stores[0], inline_agent=True))
    sc05 = Sc05Process(stores[4], Sc04Process(stores[3], Sc03Process(stores[2], sc02)))
    from app.sc06.process import Sc06Process

    return Sc06Process(
        SqliteSc01Store(sqlite3.connect(":memory:", check_same_thread=False), "sc06"), sc05
    )._finish_sc05(None, "embedded")


@pytest.fixture(scope="module")
def doc(sc05_doc: dict) -> dict:
    from datetime import UTC, datetime

    return domain.new_review(SEED, sc05_doc["demo_run_id"], sc05_doc, created_at=datetime.now(UTC))


def calc_of(doc: dict) -> dict:
    return doc["calculations"][doc["report"]["calculation_version"]]


def rows_with(sheet: str, change) -> dict[str, list[dict]]:
    """시드 시트의 복제본에 변경을 넣는다 — 원 시드는 그대로다."""
    rows = copy.deepcopy(SEED.rows(sheet))
    change(rows)
    return {sheet: rows}


def review_with(sc05_doc: dict, rows: dict[str, list[dict]]) -> dict:
    from datetime import UTC, datetime

    return domain.new_review(
        SEED,
        sc05_doc["demo_run_id"],
        sc05_doc,
        created_at=datetime.now(UTC),
        batch=domain.load_batch(SEED, rows),
    )


def _won(value) -> int:
    return int(Decimal(str(value)))


# ─────────────────────────── 원천 보존 ───────────────────────────


def test_v06_keeps_the_71_earlier_sheets_identical_to_v05() -> None:
    """AT61-12 — 기존 71개 시트는 v0.5와 한 칸도 다르지 않다. 장면 6 시트만 더해졌다."""
    old = load_workbook(SC05_PATH, data_only=True, read_only=True)
    new = load_workbook(SEED_PATH, data_only=True, read_only=True)
    try:
        assert new.sheetnames[: len(old.sheetnames)] == old.sheetnames
        assert len(old.sheetnames) == 71
        assert all(n.startswith("SC06") for n in new.sheetnames[71:])
        for name in old.sheetnames:
            assert list(old[name].iter_rows(values_only=True)) == list(
                new[name].iter_rows(values_only=True)
            ), name
    finally:
        old.close()
        new.close()


def test_baseline_snapshot_equals_the_sc05_fc51_run(doc: dict, sc05_doc: dict) -> None:
    """CK61-14 · CT601 — BASE61의 기대는 장면 5 FC51 산정본의 같은 과제·월 효과와 원 단위로 같다."""
    assert doc["sc05"]["official"] == "FC51" and doc["sc05"]["handoff_id"] == "HANDOFF51"
    assert calc_of(doc)["snapshot_residual"] == 0
    for b in doc["baselines"]:
        assert calc.expected_net(b) == b["expected_sheet"], b["baseline_id"]


def test_fc51_mismatch_refuses_to_load(sc05_doc: dict) -> None:
    """기준을 실적에 맞춰 고치지 않는다 — FC51과 다른 장면 5 결과면 적재를 막는다(409)."""
    tampered = copy.deepcopy(sc05_doc)
    run = tampered["runs"][tampered["forecasts"]["FC51"]["run_id"]]
    for e in run["effects"]:
        if e["task_id"] == "ACT42" and e["month"] == "2027-11-01":
            e["net"] = str(int(e["net"]) - 1_200_000)
    with pytest.raises(DomainError) as err:
        review_with(tampered, {})
    assert err.value.status == 409


# ─────────────────────────── 파생 시트 대조 ───────────────────────────


def test_task_performance_matches_sheet(doc: dict) -> None:
    """
    U601-03 · AT61-02 — 3과제 기대·실현·차이·효과 판단. TOTAL은 과제 계산에서 낸다(다시 합산하지
    않는다).
    """
    c = calc_of(doc)
    tasks = {t["task_id"]: t for t in domain.summary(doc, actor="u-fin")["tasks"]}
    for row in SEED.rows("SC06과제성과"):
        if row["task_id"] == "TOTAL":
            assert c["totals"]["expected_krw"] == _won(row["분기기대"])
            assert c["totals"]["actual_krw"] == _won(row["분기순효과"])
            assert c["totals"]["variance_krw"] == _won(row["차이"])
            continue
        t = tasks[row["task_id"]]
        assert (t["expected_krw"], t["actual_krw"], t["variance_krw"]) == (
            _won(row["분기기대"]),
            _won(row["분기순효과"]),
            _won(row["차이"]),
        )
        assert t["performance"] == row["효과판단"]
        assert t["activity_judgment"] == "COMPLETION_CONFIRMED" and row["활동판단"] == "완료 확인"
        assert t["completed_at"] == kst(row["완료일"])
    assert "TOTAL" not in doc["tasks"]


def test_monthly_effects_match_sheet(doc: dict) -> None:
    """
    SC06월효과 9행 전부 — 기준·실제·운영비·고정비 절감 0·확인 효과·차이. ACT42 일회성비는 Q4에 다시
    빼지 않는다.
    """
    months = {m["effect_id"]: m for m in calc_of(doc)["months"]}
    assert len(months) == 9
    for row in SEED.rows("SC06월효과"):
        m = months[row["effect_id"]]
        assert m["month"] == kst(row["월"])[:10]
        assert m["expected"] == _won(row["기대순효과"])
        assert m["qty"] == _won(row["적용수량"])
        assert m["baseline_or_avoided"] == _won(row["기준비용또는회피비"])
        assert m["actual_cost"] == _won(row["실제비용"])
        assert m["operating_cost"] == _won(row["추가운영비"])
        assert m["fixed_saving"] == _won(row["고정비절감"]) == 0
        assert m["net"] == _won(row["확인순효과"])
        assert m["variance"] == _won(row["차이"])


def test_cause_groups_match_sheet_and_reconcile(doc: dict) -> None:
    """
    AT61-03 · M607 — INTERNAL 4 / SUPPLY 9 / MIXED 1, 비용 합계 16.8백만원 = 긴급 14건 전체. 대사
    잔차 0.
    """
    causes = calc_of(doc)["causes"]
    groups = {g["cause"]: g for g in causes["groups"]}
    for row in SEED.rows("SC06원인집계"):
        if row["원인코드"] == "TOTAL":
            assert causes["total"] == {
                "count": row["건수"],
                "actual_urgent_cost_krw": _won(row["긴급비용"]),
            }
            continue
        g = groups[row["원인코드"]]
        assert (g["count"], g["actual_urgent_cost_krw"]) == (row["건수"], _won(row["긴급비용"]))
    assert "UNKNOWN" not in groups
    assert (
        causes["count_residual"] == 0 and causes["cost_residual"] == 0 and causes["problems"] == []
    )


def test_operating_metrics_match_sheet(doc: dict) -> None:
    """SC06운영지표 — 주문 180·긴급 14·확인 지연 5(복합 포함)·공유 지연 0·납기·품질 이상 0."""
    metrics = {m["metric_id"]: m for m in calc_of(doc)["metrics"]}
    for row in SEED.rows("SC06운영지표"):
        got = metrics[row["metric_id"]]["value"]
        if row["metric_id"] == "M61-03":
            assert Decimal(got) == Decimal(str(row["값"])).quantize(Decimal("0.00000001"))
        else:
            assert got == _won(row["값"]), row["metric_id"]
    assert metrics["M61-07"]["unconfirmed"] == 0


def test_quarter_profit_matches_sheet(doc: dict) -> None:
    """
    SC06분기손익 — FC51 관리손익에 효과 차이만 더한다. Q4 902,868,880(목표 대비 +31,878,880). 연간
    FC51은 그대로다.
    """
    q = calc_of(doc)["quarter_profit"]
    rows = SEED.rows("SC06분기손익")
    for got, row in zip([*q["months"], q["quarter"]], rows, strict=True):
        assert (got["target"], got["fc51_profit"], got["fc51_effect"]) == (
            _won(row["Target"]),
            _won(row["FC51관리손익"]),
            _won(row["FC51과제효과"]),
        )
        assert (got["confirmed_effect"], got["effect_diff"], got["profit"], got["vs_target"]) == (
            _won(row["확인과제효과"]),
            _won(row["효과차이"]),
            _won(row["가상분기실적"]),
            _won(row["목표차이_실적빼기목표"]),
        )
    assert doc["sc05"]["annual"] == "3428986680" and doc["sc05"]["target"] == "3477290000"
    assert doc["sc05"]["gap"] == "48303320"


def test_seventeen_checks_are_recomputed(doc: dict) -> None:
    """SC06검증 17항목 — PASS 글을 옮기지 않고 계산값을 기대값과 대조한다."""
    checks = {c["check_id"]: c["value"] for c in calc_of(doc)["checks"]}
    rows = SEED.rows("SC06검증")
    assert len(rows) == len(checks) == 17
    for row in rows:
        assert checks[row["check_id"]] == _won(row["기대값"]), row["check_id"]


def test_summary_sheet_headline_numbers(doc: dict) -> None:
    """SC06요약의 머리 숫자가 계산과 같다(총 기대·확인·차이·ACT42·긴급 건수)."""
    s = {r["항목"]: r["값"] for r in SEED.rows("SC06요약")}
    c = calc_of(doc)
    act42 = next(t for t in c["tasks"] if t["task_id"] == "ACT42")
    assert c["totals"]["expected_krw"] == _won(s["총 기대효과"])
    assert c["totals"]["actual_krw"] == _won(s["확인된 순개선효과"])
    assert c["totals"]["variance_krw"] == _won(s["기대 대비 차이"])
    assert (act42["expected_krw"], act42["actual_krw"]) == (
        _won(s["ACT42 기대효과"]),
        _won(s["ACT42 실현효과"]),
    )
    assert c["causes"]["total"]["count"] == _won(s["ACT42 긴급구매"])


def test_report_text_numbers_match_the_bound_calculation(doc: dict) -> None:
    """
    CT611 — 확인 이력 다섯 단계가 같은 RPT61-v1·순서대로이고, 보고 결론의 숫자가 묶인 산정과 같다.
    """
    assert doc["report"]["problems"] == []
    assert (
        doc["report"]["state"] == "REPORTED" and doc["report"]["calculation_version"] == "CALC61-v1"
    )


# ─────────────────────────── 예외 (X61-*, AT61-06~12) ───────────────────────────


def _set_order(order_id: str, **fields):
    def change(rows):
        row = next(r for r in rows if r["order_id"] == order_id)
        row.update(fields)

    return change


def test_x61_01_unconfirmed_cost_holds_the_task_effect(sc05_doc: dict) -> None:
    """
    AT61-07 · X61-01 — 운송비 null → ACT42 산정 보류, 합계 None. 확인된 ACT41·43은 소계로만(확정
    합계 아님).
    """
    review = review_with(sc05_doc, rows_with("SC06구매거래", _set_order("PO61-1103", 운송비=None)))
    c = calc_of(review)
    act42 = next(t for t in c["tasks"] if t["task_id"] == "ACT42")
    assert act42["actual_krw"] is None and act42["variance_krw"] is None
    assert act42["calculation_status"] == "PENDING"
    nov = next(m for m in c["months"] if m["effect_id"] == "EF61-ACT42-11")
    assert nov["net"] is None and nov["unconfirmed_orders"] == ["PO61-1103"]
    assert next(m for m in c["months"] if m["effect_id"] == "EF61-ACT42-10")["net"] == 8_200_000
    assert c["totals"]["actual_krw"] is None and c["totals"]["calculation_status"] == "PENDING"
    assert c["totals"]["confirmed_subtotal"] == {
        "tasks": ["ACT41", "ACT43"],
        "expected_krw": 61_200_000,
        "actual_krw": 61_200_000,
    }
    assert c["data_status"] == "INCOMPLETE"
    supply = next(g for g in c["causes"]["groups"] if g["cause"] == "SUPPLY")
    assert supply["count"] == 9 and supply["actual_urgent_cost_krw"] is None
    # 확정된 보고 결론의 숫자를 이 산정으로 보일 수 없다
    assert any("보고 결론의 숫자" in p for p in review["report"]["problems"])


def test_x61_02_no_alternative_evidence_holds_the_judgment(sc05_doc: dict) -> None:
    """
    AT61-07 · X61-02 — 대체 근거도 없으면 판단 보류. 담당·보완기한·수집 상태(일부 보완 중)는 그대로
    보인다.
    """

    def drop(rows):
        next(r for r in rows if r["judgment_id"] == "JDG61")["대체근거"] = None

    review = review_with(sc05_doc, rows_with("SC06완료판단", drop))
    s = domain.summary(review, actor="u-fin")
    j = next(x for x in s["judgments"] if x["judgment_id"] == "JDG61")
    assert j["judgment_state"] == "HOLD" and j["activity_state"] == "HOLD"
    assert j["collection_state"] == "PARTIAL" and j["missing"]
    assert j["owner"]["account"] == "USR-BUY" and j["due_at"] == "2028-01-18T18:00:00+09:00"
    assert next(t for t in s["tasks"] if t["task_id"] == "ACT42")["activity_judgment"] == "HOLD"


def test_x61_03_duplicate_order_is_held_not_double_counted(sc05_doc: dict) -> None:
    """AT61-08 · X61-03 — 같은 order_id가 또 오면 뒤의 행은 대사 대기. 효과·건수는 그대로다."""

    def dup(rows):
        rows.append(dict(next(r for r in rows if r["order_id"] == "PO61-1103")))

    review = review_with(sc05_doc, rows_with("SC06구매거래", dup))
    c = calc_of(review)
    assert c["totals"]["actual_krw"] == 87_000_000 and c["causes"]["total"]["count"] == 14
    assert [o["order_id"] for o in review["held_orders"]] == ["PO61-1103"]
    assert any("중복" in p for p in review["batch"]["problems"])
    assert c["data_status"] == "INCOMPLETE"


def test_x61_04_agent_draft_without_human_review_stays_a_candidate(sc05_doc: dict) -> None:
    """
    AT61-08 · X61-04 — 사람 보정 전 AG62는 후보다. 원인 집계는 거래의 최종원인(사람 확인)으로만
    한다.
    """

    def blank(rows):
        r = next(x for x in rows if x["run_id"] == "AG62")
        r["사람 보정·판단"] = None
        r["확인자"] = None
        r["확인_KST"] = None

    review = review_with(sc05_doc, rows_with("SC06Agent", blank))
    v = domain.variance(review, "ACT42", actor="u-fin")
    assert v["agent"]["state"] == "CANDIDATE" and v["agent"]["candidate_cause"] == "INTERNAL"
    assert v["agent"]["human"] is None and v["order"]["cause"] == "SUPPLY"
    assert {g["cause"]: g["count"] for g in v["cause_groups"]} == {
        "INTERNAL": 4,
        "SUPPLY": 9,
        "MIXED": 1,
    }


def test_at61_06_and_x61_05_changed_cost_makes_a_new_calculation_version(doc: dict) -> None:
    """
    AT61-06 · X61-05 — 확정 뒤 PO61-1103 운송비 +100,000 → 새 산정 CALC61-v2(ACT42 25,700,000 ·
    합계 86,900,000).
    RPT61-v1과 그 숫자는 그대로이고, ACT42가 걸린 확인(CF61-B)부터 종합·확정·보고를 다시 해야 한다.
    ACT41 확인은 아니다.
    """
    review = copy.deepcopy(doc)
    domain.reimport(
        review,
        SEED,
        rows_with("SC06구매거래", _set_order("PO61-1103", 운송비=1_000_000)),
        actor="u-fin",
        reason="운송비 정산서 재확인",
    )
    assert (
        review["current_calculation"] == "CALC61-v2"
        and review["report"]["calculation_version"] == "CALC61-v1"
    )
    new = review["calculations"]["CALC61-v2"]
    assert next(t for t in new["tasks"] if t["task_id"] == "ACT42")["actual_krw"] == 25_700_000
    assert new["totals"]["actual_krw"] == 86_900_000
    assert new["changed_tasks"] == ["ACT42"]
    assert new["reconfirm"] == ["CF61-B", "CF61-F", "CF61-R", "CF61-M"]
    s = domain.summary(review, actor="u-fin")
    assert s["totals"]["actual_krw"] == 87_000_000 and s["calculation_version"] == "CALC61-v1"
    assert s["newer_calculation"]["calculation_version"] == "CALC61-v2"
    assert s["newer_calculation"]["totals"]["actual_krw"] == 86_900_000
    # 원본 입력의 재적재는 새 버전을 만들지 않는다 — 같은 지문의 CALC61-v1이 다시 현재다
    domain.reimport(review, SEED, {}, actor="u-fin", reason="원본 재확인")
    assert sorted(review["calculations"]) == ["CALC61-v1", "CALC61-v2"]
    assert review["current_calculation"] == "CALC61-v1"
    assert domain.summary(review, actor="u-fin")["newer_calculation"] is None


def test_x61_06_proposal_effect_and_budget_are_not_estimated(doc: dict) -> None:
    """
    AT61-10 · X61-06 — PROP61 효과·예산은 null + 미산정(0원 승인 아님). FUP61 예산 0은 확인된 0이다.
    """
    f = doc["followups"]
    assert f["PROP61"]["effect"] == {
        "krw": None,
        "estimation_status": "NOT_ESTIMATED",
        "note": "NOT_ESTIMATED",
    }
    assert (
        f["PROP61"]["budget"]["krw"] is None
        and f["PROP61"]["budget"]["estimation_status"] == "NOT_ESTIMATED"
    )
    assert f["PROP61"]["state"] == "PROPOSED"
    assert (
        f["FUP61"]["budget"]["krw"] == 0
        and f["FUP61"]["budget"]["estimation_status"] == "CONFIRMED"
    )
    assert (
        f["FUP61"]["effect"]["krw"] is None
        and f["FUP61"]["effect"]["estimation_status"] == "NOT_ESTIMATED"
    )
    assert f["FUP62"]["effect"]["estimation_status"] == "NOT_APPLICABLE"


def test_x61_07_shortfall_added_to_cause_costs_is_flagged() -> None:
    """
    AT61-11 · X61-07 — 원인별 비용에 기대 미달(4.8백만원)을 더한 표는 대사 오류다. 서로 다른 지표로
    남긴다.
    """
    groups = [
        {"cause": "INTERNAL", "count": 4, "actual_urgent_cost_krw": 4_800_000 + 4_800_000},
        {"cause": "SUPPLY", "count": 9, "actual_urgent_cost_krw": 10_800_000},
        {"cause": "MIXED", "count": 1, "actual_urgent_cost_krw": 1_200_000},
    ]
    out = calc.reconcile_causes(groups, 14, 16_800_000, variance=-4_800_000)
    assert out["cost_residual"] == 4_800_000
    assert any("X61-07" in p for p in out["problems"])


def test_x61_08_partial_collection_with_alternatives_is_reviewable(doc: dict) -> None:
    """
    AT61-05 · AT61-11 · X61-08 — JDG61: 일부 보완 중 + 대체 근거 3건 → 검토 가능. 미비·담당·기한은
    지우지 않는다.
    """
    j = next(
        x for x in domain.summary(doc, actor="u-fin")["judgments"] if x["judgment_id"] == "JDG61"
    )
    assert j["collection_state"] == "PARTIAL" and j["judgment_state"] == "REVIEWABLE"
    assert j["alternatives"] == [
        "DOC59 전자확인",
        "DOC51 운영기준",
        "DOC61-HAND 적용·대행 운영기록",
    ]
    assert (
        j["missing"]
        and j["judged_by"]["account"] == "EX-BUY"
        and j["due_at"] == "2028-01-18T18:00:00+09:00"
    )


def test_load_rejects_malformed_amounts_and_reads_comma_strings(sc05_doc: dict) -> None:
    """
    구현명세 8장 1단계 — 쉼표 문자열은 읽고, 오류값은 적재 전체를 거절한다(일부 적재 상태를
    내보내지 않는다).
    """
    ok = review_with(sc05_doc, rows_with("SC06구매거래", _set_order("PO61-1103", 운송비="900,000")))
    assert calc_of(ok)["totals"]["actual_krw"] == 87_000_000
    with pytest.raises(DomainError) as err:
        domain.load_batch(
            SEED, rows_with("SC06구매거래", _set_order("PO61-1103", 운송비="9OO,000"))
        )
    assert err.value.status == 422


def test_urgent_order_missing_links_or_approved_after_order_waits(sc05_doc: dict) -> None:
    """
    구현명세 8장 2·3단계 — 긴급 주문의 승인자 누락·승인 전 발주는 대사 대기. 그 달 효과는 산정
    보류다.
    """
    review = review_with(
        sc05_doc,
        rows_with(
            "SC06구매거래",
            _set_order("PO61-1201", 긴급승인자=None, 긴급승인_KST="2027-12-06 13:00:00"),
        ),
    )
    o = review["orders"]["PO61-1201"]
    assert o["held"] and len(o["hold_reasons"]) == 2
    dec = next(m for m in calc_of(review)["months"] if m["effect_id"] == "EF61-ACT42-12")
    assert dec["net"] is None


def test_fixed_saving_is_rejected(sc05_doc: dict) -> None:
    """CT605 — 외주 확대에 따른 잔존 고정비를 절감으로 만들지 않는다. 0이 아니면 산정 거절이다."""

    def saving(rows):
        rows[0]["고정비절감"] = 1_000_000

    with pytest.raises(DomainError) as err:
        review_with(sc05_doc, rows_with("SC06실행실적", saving))
    assert err.value.status == 422


def test_at61_12_proposal_scope_does_not_overlap_act43(doc: dict) -> None:
    """
    AT61-12 — Q4 주문은 전부 A-b/P-COM01(ACT42 근거)이다. PROP61 범위 판단이 ACT43 P-MAT02와 다름을
    남긴다.
    """
    assert {(o["model_id"], o["part_id"]) for o in doc["orders"].values()} == {("A-b", "P-COM01")}
    assert "P-MAT02" in doc["followups"]["PROP61"]["scope_check"]
    assert "P-MAT02" in doc["evidence"]["DOC61-ALT"]["body"]


def test_at61_14_followups_and_next_report(doc: dict) -> None:
    """AT61-14 — FUP61 1/21 · PROP61 2/8 · FUP62 1/18 · 다음 보고 2/11 · RISK61 OPEN."""
    f = doc["followups"]
    assert (
        f["FUP61"]["due_at"] == "2028-01-21T18:00:00+09:00" and f["FUP61"]["state"] == "IN_PROGRESS"
    )
    assert f["PROP61"]["due_at"] == "2028-02-08T18:00:00+09:00"
    assert f["FUP62"]["due_at"] == "2028-01-18T18:00:00+09:00"
    assert f["RISK61"]["state"] == "OPEN" and doc["risk"]["state"] == "OPEN"
    assert doc["report"]["next_review"] == "2028-02-11T10:00:00+09:00"
