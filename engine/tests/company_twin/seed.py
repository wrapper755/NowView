"""
회사 시스템 쌍둥이의 기준 자료 — 엔진 시험 전용. 회사 시스템의 원본은 autopart1 리포의
기간계(`legacy/`)이고, 이것은 그 HTTP 계약(`docs/worker-interfaces.md` §4.2)을 같은 프로세스에서
흉내 내 엔진 시험이 회사 API를 지나게 한다. 자료는 엔진 엑셀(`engine/seed/SYN-SC01-v0.2.xlsx`)의
원천 시트와 엔진 명단(`domain.USERS`)에서 읽는다 — 기간계의 사본과 같은 내용이다.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from functools import cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKBOOK = Path(__file__).resolve().parents[2] / "seed" / "SYN-SC01-v0.2.xlsx"


def _as_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def _rows(ws: Any, first_key: str) -> Iterator[dict[str, Any]]:
    """`first_key`로 시작하는 헤더 행을 찾고, 그 아래 행들을 열 이름 → 값으로 돌려준다."""
    header: list[str] | None = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == first_key:
                header = [str(c) if c is not None else "" for c in row]
            continue
        if not any(v is not None for v in row):
            continue
        if row[0] == "rule_key":  # 버전기준 시트 안의 두 번째 표(NowView 규칙) — 읽지 않는다
            return
        yield dict(zip(header, row, strict=False))
    if header is None:
        raise ValueError(f"{ws.title}: '{first_key}' 헤더 행이 없다")


def _key(r: dict[str, Any]) -> dict[str, str]:
    return {
        "customer_id": r["customer_id"],
        "plant_id": r["plant_id"],
        "model_id": r["model_id"],
        "period": _as_date(r["period"]).isoformat(),
    }


def _plain(v: Any) -> Any:
    if isinstance(v, datetime | date):
        return v.isoformat()
    return v


@dataclass(frozen=True)
class CompanyData:
    """회사의 불변 기록 — JSON으로 바로 내보낼 모양이다(금액은 문자열 DECIMAL)."""

    users: list[dict[str, str]]
    versions: list[dict[str, Any]]
    plans: dict[str, list[dict[str, Any]]]  # T01 · F01 → 행
    customer_plans: dict[str, list[dict[str, Any]]]  # CP01 · CP02 … → 원천 행
    actuals: list[dict[str, Any]]
    costs: dict[str, list[dict[str, Any]]]  # C01 → 행
    fixed_costs: dict[str, list[dict[str, Any]]]  # FC01 → 행

    @property
    def models(self) -> list[str]:
        return sorted({r["model_id"] for r in self.plans["T01"]})


#: 회사가 가진 버전 — 엑셀 버전기준 시트에는 NowView가 계산할 검토 전망 F02(기대값)·규칙
#: R01-v1과 기준 시점 뒤에 도착할 원천 CP02도 있다. 회사의 확정 기록이 아니라서 내보내지 않는다.
COMPANY_VERSIONS = ("T01", "CP01", "F01", "C01")


@cache
def load() -> CompanyData:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    versions = [
        {k: _plain(v) for k, v in r.items()}
        for r in _rows(wb["버전기준"], "version_id")
        if r["version_id"] in COMPANY_VERSIONS
    ]

    plans: dict[str, list[dict[str, Any]]] = {}
    for r in _rows(wb["물량스냅샷"], "version_id"):
        if r["version_id"] not in COMPANY_VERSIONS:
            continue
        plans.setdefault(r["version_id"], []).append({**_key(r), "quantity": int(r["quantity"])})

    customer_plans: dict[str, list[dict[str, Any]]] = {}
    for r in _rows(wb["고객원천"], "source_id"):
        q = r["quantity"]
        customer_plans.setdefault(r["customer_plan_id"], []).append(
            {
                "source_id": r["source_id"],
                "customer_plan_id": r["customer_plan_id"],
                **_key(r),
                "quantity": None if q is None else int(q),  # 누락은 0이 아니다(G13)
                "unit": r["unit"],
                "quantity_scope": r["quantity_scope"],
                "allocation_scope": r["allocation_scope"],
            }
        )

    actuals = [
        {
            **_key(r),
            "quantity": int(r["quantity"]),
            "cutoff_date": _as_date(r["cutoff_date"]).isoformat(),
        }
        for r in _rows(wb["실적"], "actual_version")
    ]

    costs: dict[str, list[dict[str, Any]]] = {}
    for r in _rows(wb["단위원가"], "cost_version"):
        costs.setdefault(str(r["cost_version"]), []).append(
            {
                **_key(r),
                "unit_price": str(r["unit_price"]),
                "unit_variable_cost": str(r["unit_variable_cost"]),
                "unit_cm": str(r["unit_cm"]),
            }
        )

    fixed: list[dict[str, Any]] = [
        {
            "plant_id": r["plant_id"],
            "period": _as_date(r["period"]).isoformat(),
            "total_fixed": str(r["total_fixed"]),
        }
        for r in _rows(wb["고정비"], "plant_id")
    ]
    wb.close()
    from app.sc01.domain import USERS

    org = {
        "users": [
            {"user_id": uid, **u} for uid, u in USERS.items() if u["role"] != "시연관리자"
        ]
    }
    return CompanyData(
        users=org["users"],
        versions=versions,
        plans=plans,
        customer_plans=customer_plans,
        actuals=actuals,
        costs=costs,
        fixed_costs={"FC01": fixed},
    )


#: 고객이 보낼 수 있는 원천 — 시연 사례. CP02(기본)와 그 변형(EX01·02·04·05), 다른 계획 CP03.
#: 값은 CP02 원천 행에서 (모델, 월)의 물량만 바꾼다. None은 누락이다.
FIXTURES: dict[str, tuple[str, dict[tuple[str, int], int | None]]] = {
    "CP02": ("CP02", {}),
    "EX01": ("CP02", {("A-b", 7): 9500, ("A-c", 7): 5900}),
    "EX02": ("CP02", {("A-b", 7): None}),
    "CP03": ("CP03", {("A-b", 7): 9000}),
    "EX04": ("CP02", {("A-c", 7): 6500}),
    "EX05": ("CP02", {("A-b", 7): 0}),
}

#: 기준 시점(S0)에 이미 받아 확인한 원천 — 수신함이 아니라 회사 기록이다.
BASELINE_PLANS = ("CP01",)


def fixture_rows(fixture_id: str) -> tuple[str, list[dict[str, Any]]]:
    plan_id, overrides = FIXTURES[fixture_id]
    rows = []
    for row in load().customer_plans["CP02"]:
        month = date.fromisoformat(row["period"]).month
        q = overrides.get((row["model_id"], month), row["quantity"])
        rows.append({**row, "customer_plan_id": plan_id, "quantity": q})
    return plan_id, rows
