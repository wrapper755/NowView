"""
SC02 규칙 계산 — **통합 가상 데이터 v0.2의 파생 시트와 대조한다**(TEST-22).

시드는 원천 시트만 읽는다(`app/sc02/seed.py`). 여기서는 파생 시트(단가안비교·통합손익·정합성검증·
예외데이터)를 따로 읽어 서비스가 다시 계산한 값과 맞춰 본다 — 숫자를 테스트에 옮겨 적지 않는다.
"""

from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from functools import cache
from typing import Any

import pytest
from openpyxl import load_workbook

from app.sc01.calc import Risk
from app.sc01.seed import load_seed as load_sc01_seed
from app.sc02 import calc
from app.sc02.calc import Option, analyze
from app.sc02.seed import SEED_PATH, MonthRow, load_seed

RISK = {"정상": "NORMAL", "관심": "ATTENTION", "중요": "CRITICAL"}


@cache
def sheet(name: str, first: str) -> list[dict[str, Any]]:
    wb = load_workbook(SEED_PATH, data_only=True, read_only=True)
    out: list[dict[str, Any]] = []
    header = None
    for row in wb[name].iter_rows(values_only=True):
        if header is None:
            if row and row[0] == first:
                header = list(row)
            continue
        if row and row[0] is not None:
            out.append(dict(zip(header, row, strict=False)))
    wb.close()
    return out


def option(version: str, price: int | None, effective: date = date(2027, 7, 1)) -> Option:
    seed = load_seed()
    return Option(
        version,
        version,
        "P-COM01",
        "SUP-01",
        seed.setting_decimal("old_price"),
        None if price is None else Decimal(price),
        effective,
    )


def request() -> Option:
    return option("PV-REQ01", int(load_seed().settings["request_price"]))


# ─────────────────────────── 원천 ───────────────────────────


def test_2027_rows_are_the_same_as_the_sc01_seed() -> None:
    """SC02-10 — 장면 1의 T01·F01·F02·C01을 재사용한다. 두 시드가 갈라지면 장면 연결이 틀린다."""
    s1, s2 = load_sc01_seed(), load_seed()
    for r in s2.rows:
        if r.period.year != 2027:
            continue
        key = ("CUS-HK", "ULS", r.model_id, r.period)
        assert r.t01 == s1.plans["T01"][key]
        assert r.f01 == s1.plans["F01"][key]
        assert r.f02 == s1.plans["F02"][key], key
        assert r.unit_cm == s1.costs[key].cm
        assert s2.fixed[r.period] == s1.fixed_costs[("ULS", r.period)]


def test_2028_is_the_internal_assumption_f01_times_97_percent() -> None:
    """DATA-21 — 2028-01~07은 2027 동일월 F01 × 97% 정수 반올림. 7월 b의 −700을 반복하지 않는다."""
    seed = load_seed()
    f01_2027 = {(r.period.month, r.model_id): r.f01 for r in seed.rows if r.period.year == 2027}
    factor = Decimal(str(seed.settings["next_year_factor"]))
    for r in seed.rows:
        if r.period.year == 2028:
            # 엑셀 ROUND는 0.5를 올린다 — Decimal 기본(짝수 반올림)이면 4898.5가 4898이 된다.
            raw = Decimal(f01_2027[(r.period.month, r.model_id)]) * factor
            expected = int(raw.to_integral_value(rounding=ROUND_HALF_UP))
            assert r.f02 == r.f01 == expected
            assert r.t01 is None and r.assumed


def test_integrity_checks_v05_v06_v09_v16_v17() -> None:
    expected = {r["ID"]: r["기대값"] for r in sheet("정합성검증", "ID")}
    got = calc.integrity(load_seed())
    for vid, value in got.items():
        assert value == expected[vid], vid


# ─────────────────────────── 단가안 비교 (T23·T24) ───────────────────────────


@pytest.mark.parametrize(
    "row", sheet("단가안비교", "price_version"), ids=lambda r: r["price_version"]
)
def test_price_options_match_the_comparison_sheet(row: dict[str, Any]) -> None:
    """요청·내부제시·협상·기안·승인안 — 인상률·잔여/12개월 수량·증가액·위험·승인선·부담 감소."""
    a = analyze(load_seed(), option(row["price_version"], row["단가"]), base=request())
    assert Decimal(a["rate"]) == Decimal(str(row["인상률"]))
    assert a["remaining"]["purchase_qty"] == row["잔여구매량"]
    assert a["next12"]["purchase_qty"] == row["12개월구매량"]
    assert Decimal(a["remaining"]["increase"]) == row["잔여구매증가"]
    assert Decimal(a["next12"]["increase"]) == row["12개월구매증가"]
    assert a["rate_risk"] == RISK[row["비율위험"]]
    assert a["amount_risk"] == RISK[row["금액위험"]]
    assert a["risk"] == RISK[row["종합위험"]]
    assert a["ceo_required"] is (row["승인선미리보기"] == "대표 대상")
    assert Decimal(a["comparison"]["reduction"]) == row["요청대비부담감소"]
    # 12개월 CM 영향 = −구매 증가 — 시차·재고 0의 단순화다(CALC-24). 필드는 따로 둔다.
    next12 = next(p for p in a["periods"] if p["kind"] == "next12")
    assert Decimal(next12["cost_effect"]) == row["12개월CM영향"]


# ─────────────────────────── 손익 연결 (T27·T28) ───────────────────────────


@pytest.mark.parametrize("row", sheet("통합손익", "기간")[:17], ids=lambda r: r["기간"])
def test_combined_pnl_matches_every_period_row(row: dict[str, Any]) -> None:
    """2027 월별 · Q3 · Q4 · 잔여 · 연간 · 향후 12개월 — 기준 → 물량 → 단가 → 통합, Target 부족."""
    a = analyze(load_seed(), option("PV-APP01", 2120))
    p = next(x for x in a["periods"] if x["label"] == row["기간"])
    assert p["computable"]
    for key, col in [
        ("revenue_f02", "F02매출"),
        ("fixed", "공장고정비"),
        ("profit_f01", "F01기준손익"),
        ("profit_volume", "물량반영손익"),
        ("volume_effect", "장면1물량영향"),
        ("cost_effect", "장면2원가영향"),
        ("profit_combined", "통합손익"),
        ("total_change", "전체변동"),
        ("bridge_diff", "브리지차이"),
    ]:
        assert Decimal(p[key]) == row[col], (row["기간"], key)
    if row["Target손익"] is None:
        # 2028 Target 미설정 — 12개월 Target 차이는 미산출(CALC-23). 0이 아니다.
        assert p["target_profit"] is None and p["target_gap"] is None
    else:
        assert Decimal(p["target_profit"]) == row["Target손익"]
        assert Decimal(p["target_gap"]) == row["Target부족"]


def test_q3_by_model_keeps_c_to_e_unchanged() -> None:
    """T28·V08 — Q3 모델별. c~e 직접 단가영향 0. 고정비는 모델에 배부하지 않는다."""
    rows = [r for r in sheet("통합손익", "2027 Q3 모델")]
    a = analyze(load_seed(), option("PV-APP01", 2120))
    got = {m["model_id"]: m for m in a["q3_models"]}
    for r in rows:
        m = got[r["2027 Q3 모델"]]
        assert Decimal(m["cm_f01"]) == r["F01 공헌이익"]
        assert Decimal(m["volume_effect"]) == r["물량영향"]
        assert Decimal(m["cost_effect"]) == r["원가영향"]
        assert Decimal(m["cm_combined"]) == r["통합 공헌이익"]
    assert all(got[m]["cost_effect"] == "0" for m in ("A-c", "A-d", "A-e"))


def test_volume_and_cost_effects_add_up_without_double_counting() -> None:
    """AT24 · CALC-22 — 장면 1 물량효과 + 장면 2 추가 원가효과 = 기준 대비 통합효과."""
    a = analyze(load_seed(), option("PV-APP01", 2120))
    year = next(p for p in a["periods"] if p["kind"] == "year")
    assert year["volume_effect"] == "-14000000"
    assert year["cost_effect"] == "-9364320"
    assert Decimal(year["volume_effect"]) + Decimal(year["cost_effect"]) == Decimal(
        year["total_change"]
    )
    # 12개월의 물량효과 −1,400만원도 월별 재집계 결과다 — 연간값을 복사한 것이 아니다(CALC-23).
    next12 = next(p for p in a["periods"] if p["kind"] == "next12")
    assert next12["volume_effect"] == "-14000000"
    assert next12["profit_combined"] == "3173163560"


def test_purchase_increase_and_cm_loss_are_kept_as_separate_fields() -> None:
    """CALC-24 — 시차·재고 0이라 절대값이 같을 뿐이다(V14·V15). 필드는 분리한다."""
    a = analyze(load_seed(), option("PV-APP01", 2120))
    remaining = next(p for p in a["periods"] if p["kind"] == "remaining")
    assert Decimal(a["remaining"]["increase"]) + Decimal(remaining["cost_effect"]) == 0
    assert "increase" in a["remaining"] and "cost_effect" in remaining


def test_cost_replaces_only_the_common_part() -> None:
    """T26 · COST-01 — a 75,000 → 75,120, b 95,000 → 95,120, c~e 그대로. 통째로 더하지 않는다."""
    rows = {
        r["model_id"]: r
        for r in calc.new_outsource(load_seed(), "P-COM01", Decimal(2000), Decimal(2120))
    }
    assert rows["A-a"]["outsource_after"] == "75120"
    assert rows["A-b"]["outsource_after"] == "95120"
    for m in ("A-c", "A-d", "A-e"):
        assert rows[m]["outsource_after"] == rows[m]["outsource_before"]


# ─────────────────────────── 예외·경계 (X02~X10) ───────────────────────────


def test_x02_proposal_at_2050_is_not_a_ceo_case() -> None:
    """X02 — 12개월 증가 7,840,600원, 금액위험 관심, 대표 비대상."""
    x02 = next(r for r in sheet("예외데이터", "case_id") if r["case_id"] == "X02")
    a = analyze(load_seed(), option("PV-PROP01", 2050))
    assert Decimal(a["next12"]["increase"]) == x02["연구매증가"]
    assert Decimal(a["rate"]) == Decimal(str(x02["인상률"]))
    assert a["amount_risk"] == RISK[x02["계산 위험"]]
    assert a["ceo_required"] is False


@pytest.mark.parametrize(
    "case",
    [r for r in sheet("예외데이터", "case_id") if r["case_id"].startswith(("X03", "X08"))],
    ids=lambda r: r["case_id"],
)
def test_x03_x08_boundaries_are_inclusive(case: dict[str, Any]) -> None:
    """X03·X08 — 규칙 fixture. 원 단위로 정확히 비교한다. 구매수량을 역산해 만들지 않는다."""
    seed = load_seed()
    rate = Decimal(str(case["인상률"]))
    amount = Decimal(case["연구매증가"])
    old = Decimal(2000)
    levels = [
        calc.rate_level(old, old + old * rate, seed.settings),
        calc.amount_level(amount, seed.settings),
    ]
    assert calc.combined_level(levels).value == RISK[case["계산 위험"]]
    assert calc.ceo_required(amount, seed.settings) is (case["계산 승인선"] == "대표 대상")


def test_rate_boundary_compares_without_division() -> None:
    """CALC-25 — 5%·10% 경계 직하값은 아래 등급이다."""
    s = load_seed().settings
    assert calc.rate_level(Decimal(2000), Decimal(2100), s) is Risk.ATTENTION
    assert calc.rate_level(Decimal(2000), Decimal("2099.9999"), s) is Risk.NORMAL
    assert calc.rate_level(Decimal(2000), Decimal(2200), s) is Risk.CRITICAL
    assert calc.rate_level(Decimal(2000), Decimal("2199.9999"), s) is Risk.ATTENTION
    assert calc.rate_level(Decimal(0), Decimal(10), s) is None


def test_x04_effective_august_uses_aug_to_next_july() -> None:
    """X04 · PERIOD-02·03 — 8/1 적용은 [2027-08, 2028-08). 협상효과는 요청안 12개월로 잰다."""
    a = analyze(load_seed(), option("PV-NEG02", 2120, date(2027, 8, 1)), base=request())
    assert a["computable"]
    assert a["next12"]["from"] == "2027-08-01" and a["next12"]["to"] == "2028-08-01"
    assert a["remaining"]["to"] == "2028-01-01"
    cmp = a["comparison"]
    assert cmp["from"] == "2027-07-01" and cmp["to"] == "2028-07-01"
    # 가격 효과와 적용일 연기 효과를 나눠 낸다 — 둘을 더하면 같은 기간의 부담 감소다.
    assert Decimal(cmp["price_effect"]) + Decimal(cmp["deferral_effect"]) == Decimal(
        cmp["reduction"]
    )
    assert Decimal(cmp["deferral_effect"]) > 0


def test_mid_month_effective_date_blocks_the_analysis() -> None:
    """PERIOD-01 — 월중 적용일은 일별 자료가 없어 확정 분석·상신을 막는다."""
    a = analyze(load_seed(), option("PV-X", 2120, date(2027, 7, 15)))
    assert not a["computable"]
    assert "MID_MONTH_EFFECTIVE" in {p["code"] for p in a["problems"]}
    assert a["next12"]["increase"] is None and a["ceo_required"] is None


def test_x06_unresolved_price_is_not_zero() -> None:
    """X06 — 협상 미합의는 NULL. 0원으로 계산하지 않는다."""
    a = analyze(load_seed(), option("PV-NEG01", None))
    assert not a["computable"]
    assert a["next12"]["increase"] is None and a["risk"] is None
    assert "NEGOTIATION_UNRESOLVED" in {p["code"] for p in a["problems"]}


def _with(rows: list[MonthRow], **change: Any) -> list[MonthRow]:
    from dataclasses import replace

    return [
        replace(r, f02=change["f02"]) if (r.period, r.model_id) == change["key"] else r
        for r in rows
    ]


def test_x09a_missing_next_year_quantity_holds_the_12_month_result() -> None:
    """X09-A — 2028-06 A-b NULL이면 12개월 미산출·판정 보류. 대표 비대상으로 처리하지 않는다."""
    seed = load_seed()
    rows = _with(seed.rows, key=(date(2028, 6, 1), "A-b"), f02=None)
    a = analyze(seed, option("PV-REQ01", 2200), rows=rows)
    assert a["next12"]["increase"] is None
    assert a["ceo_required"] is None and a["risk"] is None
    # 당해연도 잔여기간 금액도 확정하지 않는다 — 분석 자체가 막혔다.
    assert not a["computable"]


def test_zero_quantity_is_a_real_value() -> None:
    """수량 0은 유효값이다 — 누락이 아니다."""
    seed = load_seed()
    rows = _with(seed.rows, key=(date(2028, 6, 1), "A-b"), f02=0)
    a = analyze(seed, option("PV-REQ01", 2200), rows=rows)
    assert a["computable"]
    assert a["next12"]["purchase_qty"] == 156812 - next(
        r.f02 for r in seed.rows if (r.period, r.model_id) == (date(2028, 6, 1), "A-b")
    )


def test_x10_duplicate_month_model_rows_are_an_error_not_a_sum() -> None:
    """X10 — 중복키는 오류다. SUM으로 합산하지 않는다."""
    seed = load_seed()
    dup = next(r for r in seed.rows if (r.period, r.model_id) == (date(2027, 9, 1), "A-a"))
    a = analyze(seed, option("PV-REQ01", 2200), rows=[*seed.rows, dup])
    assert not a["computable"]
    assert "DUPLICATE_KEY" in {p["code"] for p in a["problems"]}
    assert a["next12"]["increase"] is None
