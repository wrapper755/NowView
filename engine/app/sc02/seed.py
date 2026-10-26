"""
SC02 시드 — `seed/SYN-SC02-v0.2.xlsx`(통합 가상 데이터 v0.2)의 **원천 시트**를 읽는다.

SC01과 같은 원칙이다(CLAUDE.md 4.5): 원천(설정·물량원가기준·월고정비·부품모델매핑·증빙과협상)만
적재하고 파생 시트(단가안비교·통합손익·월별계산·정합성검증)는 **읽지 않는다** — 서비스가 다시
계산한 값과 대조할 기대값이다(`tests/test_sc02_calc.py`, DATA-22).

* **행 번호를 키로 쓰지 않는다.** 시트마다 헤더 행을 찾아 열 이름으로 읽는다.
* 물량원가기준은 **목록 그대로** 둔다 — (월, 모델) 키로 접으면 중복 행(X10)이 조용히 사라진다.
* 「단가버전」 시트는 시드가 아니다 — 단계 실행 때 생길 **기대값**이다(DEMO-01: D00에 미래의
  요청·협상·승인안을 미리 적재하지 않는다).
* 2027 행은 SC01 시드와 같은 값이어야 한다(SC02-10 장면 1의 F01·F02·C01 재사용) — 테스트가 본다.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SEED_PATH = Path(__file__).resolve().parents[2] / "seed" / "SYN-SC02-v0.2.xlsx"
SEED_ID = "SYN-SC02-v0.2"


@dataclass(frozen=True)
class MonthRow:
    """물량원가기준 한 행 — 공장 ULS · A제품 · 모델 · 월. 수량은 **당사 배정 대분**이다(G10)."""

    period: date
    model_id: str
    t01: int | None  # 2027만 있다. 2028 Target은 미설정(NULL)
    f01: int | None
    f02: int | None  # 2028은 EXT28-v1 내부가정(F01 동일월 × 97%)
    price: Decimal | None
    material: Decimal | None
    outsource: Decimal | None  # C01 외주비 — 공통부품 기존 단가가 **이미 들어 있다**(COST-01)
    process_var: Decimal | None
    actual: int | None
    nature: str  # 확정실적 포함 · 미래전망 · 내부가정·미승인
    source_ref: str

    @property
    def variable_cost(self) -> Decimal | None:
        parts = (self.material, self.outsource, self.process_var)
        return None if any(p is None for p in parts) else sum(parts, Decimal(0))  # type: ignore[arg-type]

    @property
    def unit_cm(self) -> Decimal | None:
        vc = self.variable_cost
        return None if vc is None or self.price is None else self.price - vc

    @property
    def assumed(self) -> bool:
        """내부가정 행(2028) — 고객 확정 계획이 아니다. 화면이 따로 표시한다."""
        return self.period.year >= 2028


@dataclass(frozen=True)
class PartLink:
    """부품-모델 매핑. c~e는 소요량 0이 **명시적으로** 있다 — 없는 것과 다르다."""

    model_id: str
    part_id: str
    supplier_id: str
    qty_per: Decimal


@dataclass(frozen=True)
class Seed:
    #: 설정 시트 key → 값 (원 단위 정수·비율·날짜). 가상 입력값이다.
    settings: dict[str, Any]
    rows: list[MonthRow]
    #: {월: 공장 월고정비} — 공장 합계에서 한 번만 차감한다(D11).
    fixed: dict[date, Decimal]
    parts: list[PartLink]
    #: 증빙과협상 — 화면을 채울 가상 공문·메모 본문(record_id → 행)
    evidence: dict[str, dict[str, Any]]

    @property
    def models(self) -> list[str]:
        return sorted({r.model_id for r in self.rows})

    def setting_decimal(self, key: str) -> Decimal:
        return Decimal(str(self.settings[key]))

    def setting_date(self, key: str) -> date:
        return _as_date(self.settings[key])

    def qty_per(self, model_id: str, part_id: str) -> Decimal | None:
        """None = 매핑 없음(미매핑 부품은 분석을 막는다, ERR-23). 0은 「쓰지 않는다」."""
        found = next(
            (p for p in self.parts if p.model_id == model_id and p.part_id == part_id), None
        )
        return None if found is None else found.qty_per


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _int(value: Any) -> int | None:
    return None if value is None or value == "" else int(value)


def _dec(value: Any) -> Decimal | None:
    return None if value is None or value == "" else Decimal(str(value))


def _rows(ws: Any, first_key: str) -> Iterator[dict[str, Any]]:
    """`first_key`로 시작하는 헤더 행을 찾고, 그 아래 행들을 열 이름 → 값으로 돌려준다."""
    header: list[str] | None = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == first_key:
                header = [str(c) if c is not None else "" for c in row]
            continue
        if not any(v is not None for v in row):
            continue
        yield dict(zip(header, row, strict=False))
    if header is None:
        raise ValueError(f"{ws.title}: '{first_key}' 헤더 행이 없다")


@cache
def load_seed(path: Path = SEED_PATH) -> Seed:
    wb = load_workbook(path, data_only=True, read_only=True)
    settings = {str(r["key"]): r["value"] for r in _rows(wb["설정"], "key")}
    rows = [
        MonthRow(
            period=_as_date(r["period"]),
            model_id=str(r["model_id"]),
            t01=_int(r["T01_qty"]),
            f01=_int(r["F01_qty"]),
            f02=_int(r["F02_qty"]),
            price=_dec(r["판매단가"]),
            material=_dec(r["재료비"]),
            outsource=_dec(r["기존외주비"]),
            process_var=_dec(r["변동가공비"]),
            actual=_int(r["실적수량"]),
            nature=str(r["기간성격"] or ""),
            source_ref=str(r["source_ref"] or ""),
        )
        for r in _rows(wb["물량원가기준"], "row_key")
    ]
    fixed = {
        _as_date(r["period"]): Decimal(str(r["fixed_cost"]))
        for r in _rows(wb["월고정비"], "period")
    }
    parts = [
        PartLink(
            model_id=str(r["model_id"]),
            part_id=str(r["part_id"]),
            supplier_id=str(r["supplier_id"]),
            qty_per=Decimal(str(r["소요량"])),
        )
        for r in _rows(wb["부품모델매핑"], "model_id")
    ]
    evidence = {str(r["record_id"]): r for r in _rows(wb["증빙과협상"], "record_id")}
    wb.close()
    return Seed(settings=settings, rows=rows, fixed=fixed, parts=parts, evidence=evidence)
