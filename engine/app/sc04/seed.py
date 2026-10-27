"""
SC04 시드 — `seed/SYN-SC04-v0.4.xlsx`(SC01_SC02_SC03_SC04 통합가상데이터 v0.4, 2026-10-23 개선효과
수정본).

**숫자 원천은 장면 2·3과 같은 시트다** — v0.4는 v0.3의 시트를 그대로 두고(검토요약만 v0.4로 갈고
원본은
「검토요약_v02」로 옮겼다) 장면 4 시트만 더했다. 그래서 물량·원가·고정비는
`app/sc02/seed.load_seed`에
이 파일 경로를 줘서 읽는다. 장면 3 기준 손익도 같은 원천에서 장면 3 계산으로 다시
낸다(`calc.baseline`).

**장면 4 시트는 두 종류다.**

* **원천** — `SC04가정`(대안별 산정 가정)과 `SC04근거`(근거 목록). 업무 중에 읽는다: 부서 제출의
  가정이
  이 표에서 오고, 효과는 계산 서비스가 이 가정으로 만든다(구현명세 1장 — 고정 숫자를 화면에 넣지
  않는다).
* **기본 시연 종료 시점의 완성 레코드** — 비교·월효과·월손익·분석결과·예산과제·조건지시·회의·이벤트·
  인계·예외·검증. 초기 상태에 적재하지 않는다(구현명세 16장 — 미래 승인·인수를 먼저 보이지 않는다).
  각본의 입력(사람이 칠 문장)과 테스트의 기대값으로만 읽는다.

엑셀 시각은 KST 벽시계 값이다(+09:00, 가장 가까운 초) — 장면 3 `kst()`를 그대로 쓴다.
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.sc02 import seed as sc02_seed
from app.sc02.seed import Seed as NumbersSeed
from app.sc03.seed import _sheet, kst

SEED_PATH = Path(__file__).resolve().parents[2] / "seed" / "SYN-SC04-v0.4.xlsx"
SEED_ID = "SYN-SC04-v0.4"
#: 시트 이름 → 헤더 첫 칸
SHEETS = {
    "SC04비교": "대안ID",
    "SC04월손익": "기간",
    "SC04가정": "대안ID",
    "SC04월효과": "row_id",
    "SC04근거": "근거ID",
    "SC04분석결과": "결과ID",
    "SC04예산과제": "과제ID",
    "SC04조건지시": "기록ID",
    "SC04회의": "기록ID",
    "SC04이벤트": "event_id",
    "SC04인계": "field",
    "SC04예외": "case_id",
    "SC04검증": "검증ID",
}

__all__ = ["SEED_ID", "SEED_PATH", "Seed", "kst", "load_seed"]


@dataclass(frozen=True)
class Seed:
    #: 숫자 — 장면 2·3과 같은 원천 시트(물량·원가·고정비·설정)
    numbers: NumbersSeed
    #: 장면 4 시트 이름 → 행 목록(헤더 → 값). 원값 그대로다.
    sheets: dict[str, list[dict[str, Any]]]
    #: `SC04분석결과` 하단의 기존 부족 원인표(모델 · T01 · F01 · C01 단위CM · 기여) — RES44의 근거
    causes: list[dict[str, Any]]

    def rows(self, sheet: str) -> list[dict[str, Any]]:
        return self.sheets[sheet]

    def by_id(self, sheet: str) -> dict[str, dict[str, Any]]:
        key = SHEETS[sheet]
        return {str(r[key]): r for r in self.sheets[sheet]}


def _causes(ws: Any) -> list[dict[str, Any]]:
    """분석결과 시트 아래쪽의 두 번째 표 — 헤더가 「모델」인 행부터 「합계」 전까지."""
    out: list[dict[str, Any]] = []
    header: list[str] | None = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == "모델":
                header = [str(c) if c is not None else "" for c in row]
            continue
        if not row or row[0] in (None, "합계"):
            break
        out.append(dict(zip(header, row, strict=False)))
    return out


@cache
def load_seed(path: Path = SEED_PATH) -> Seed:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = {name: _sheet(wb[name], first) for name, first in SHEETS.items()}
    # 분석결과 시트의 첫 표는 「모델」 행에서 끊는다 — 원인표는 따로 읽는다
    sheets["SC04분석결과"] = [
        r for r in sheets["SC04분석결과"] if str(r["결과ID"]).startswith("RES")
    ]
    causes = _causes(wb["SC04분석결과"])
    wb.close()
    return Seed(numbers=sc02_seed.load_seed(path), sheets=sheets, causes=causes)
