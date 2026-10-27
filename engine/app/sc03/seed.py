"""
SC03 시드 — `seed/SYN-SC03-v0.3.xlsx`(SC01_SC02_SC03 통합가상데이터 v0.3).

**숫자 원천은 장면 2와 같은 시트다**(설정·물량원가기준·월고정비·부품모델매핑·증빙과협상) — v0.3은 그
시트들을 그대로 두고 장면 3 시트만 더했다(통합손익 하단 Q3 모델표는 「통합모델손익」으로 떼었다).
그래서
숫자는 `app/sc02/seed.load_seed`에 이 파일 경로를 줘서 읽는다. 장면 2 시드와 같은 값인지는 테스트가
본다.

**장면 3 시트(SC03기준·사실확인·쟁점·회의·후보·과제·이벤트·인계)는 기본 시연 종료 시점의 완성
레코드다**(SC03데이터계약). 초기 상태에 적재하지 않는다 — D30에 미래의 회신·결론·인수가 보이면 안
된다
(구현명세 13장). 여기서 읽는 것은 두 가지 용도뿐이다:

* **각본의 입력** — 회신 문장·쟁점 처리방향·질의·후보 설명·과제 초안. 사람이 화면에서 칠 내용이다
  (`demo.py`). 업무 효과는 그 입력이 BPM을 지나며 만든다.
* **기대값** — 테스트가 끝난 문서와 대조한다(`tests/test_sc03_*.py`).

과제 초안(SC03과제)만은 업무 중에 읽는다 — 결론에서 포함된 후보로 경영관리가 정리하는 초안의
원문이다
(배정 전에는 고칠 수 있다, U36).

엑셀 시각은 KST 벽시계 값이다(+09:00). 직렬 날짜의 미세 오차(10:21:59.999)는 가장 가까운 초로 맞춘다
(구현명세 9장) — UTC로 오인해 9시간 옮기지 않는다.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from functools import cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.sc01.domain import KST
from app.sc02 import seed as sc02_seed
from app.sc02.seed import Seed as NumbersSeed

SEED_PATH = Path(__file__).resolve().parents[2] / "seed" / "SYN-SC03-v0.3.xlsx"
SEED_ID = "SYN-SC03-v0.3"
#: 시트 이름 → 헤더 첫 칸
SHEETS = {
    "SC03기준": "record_id",
    "SC03사실확인": "fact_id",
    "SC03쟁점": "question_id",
    "SC03회의": "record_id",
    "SC03후보": "candidate_id",
    "SC03과제": "task_id",
    "SC03이벤트": "event_id",
    "SC03인계": "field",
    "SC03예외": "scenario_id",
    "SC03검증": "check_id",
    "SC03영업시간": "date_KST",
    "통합모델손익": "2027 Q3 모델",
    "통합손익": "기간",
}


def kst(value: Any) -> str | None:
    """엑셀 시각 → ISO8601 +09:00, 가장 가까운 초로. 빈 칸은 None(미설정)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        at = value
    else:
        at = datetime.fromisoformat(str(value).replace(" KST", "").strip())
    at = (at + timedelta(microseconds=500_000)).replace(microsecond=0)
    return at.replace(tzinfo=KST).isoformat()


@dataclass(frozen=True)
class Seed:
    #: 숫자 — 장면 2 시드와 같은 원천 시트(물량·원가·고정비·설정)
    numbers: NumbersSeed
    #: 장면 3 시트 이름 → 행 목록(헤더 → 값). 원값 그대로다 — 시각은 `kst()`로 읽는다.
    sheets: dict[str, list[dict[str, Any]]]

    def rows(self, sheet: str) -> list[dict[str, Any]]:
        return self.sheets[sheet]

    def by_id(self, sheet: str) -> dict[str, dict[str, Any]]:
        key = SHEETS[sheet]
        return {str(r[key]): r for r in self.sheets[sheet]}


def _sheet(ws: Any, first: str) -> list[dict[str, Any]]:
    header: list[str] | None = None
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == first:
                header = [str(c) if c is not None else "" for c in row]
            continue
        if row and row[0] is not None:
            out.append(dict(zip(header, row, strict=False)))
    if header is None:
        raise ValueError(f"{ws.title}: '{first}' 헤더 행이 없다")
    return out


@cache
def load_seed(path: Path = SEED_PATH) -> Seed:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = {name: _sheet(wb[name], first) for name, first in SHEETS.items()}
    wb.close()
    return Seed(numbers=sc02_seed.load_seed(path), sheets=sheets)
