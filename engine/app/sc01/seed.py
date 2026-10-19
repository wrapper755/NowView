"""
SC01 시드 — `seed/SYN-SC01-v0.2.xlsx`의 **원천 시트**를 읽는다.

엑셀이 기준 시드이자 기대 결과다(DATA-01). 원천(버전·고객원천·물량스냅샷·실적·단위원가·
고정비)만 적재하고, 파생 시트(손익계산·월분기연간·SC01검증·정합성검증)는 **읽지 않는다** —
그것들은 서비스가 다시 계산한 값과 대조할 기대값이다(`tests/test_sc01_calc.py`).
파생 값을 시드로 읽으면 계산이 틀려도 화면은 맞아 보인다.

* **행 번호를 키로 쓰지 않는다**(DATA-01). 시트마다 영문 헤더 행을 찾아 열 이름으로 읽는다.
* 시드는 불변이다. 한 번 읽어 캐시하고, 데모 실행이 만드는 값은 실행 쪽에 따로 둔다.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from functools import cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SEED_PATH = Path(__file__).resolve().parents[2] / "seed" / "SYN-SC01-v0.2.xlsx"
SEED_ID = "SYN-SC01-v0.2"

#: 비교 단위 — 고객군 × 공장 × 모델 × 대상 월 (CALC-01).
Key = tuple[str, str, str, date]


@dataclass(frozen=True)
class UnitCost:
    price: Decimal
    variable_cost: Decimal
    cm: Decimal  # 단위 공헌이익 = 가격 − 변동비


@dataclass(frozen=True)
class SourceRow:
    """고객 생산계획 원천 한 행. **당사 배정 · 월 전체 물량**이다 — 40%를 다시 곱하지 않는다."""

    source_id: str
    customer_plan_id: str
    customer_id: str
    plant_id: str
    model_id: str
    period: date
    quantity: int | None  # None = 누락. 확인된 0과 구분한다(G13).
    unit: str
    quantity_scope: str  # FULL_MONTH | REMAINING
    allocation_scope: str  # ASSIGNED_SUPPLY | CUSTOMER_TOTAL

    @property
    def key(self) -> Key:
        return (self.customer_id, self.plant_id, self.model_id, self.period)


@dataclass(frozen=True)
class Seed:
    #: version_id → 메타(종류·기준시각·상태·source_id·actual_cutoff)
    versions: dict[str, dict[str, Any]]
    #: version_id → {Key: 수량}. 시드에 있는 T01·F01·F02 (F02는 대조용 기대값이다)
    plans: dict[str, dict[Key, int]]
    #: customer_plan_id → 원천 행들
    customer_plans: dict[str, list[SourceRow]]
    #: {Key: 실적 수량} — 1~5월 마감, 6월은 6/15 누계
    actuals: dict[Key, int]
    actual_cutoffs: dict[Key, date]
    #: {Key: 단가·변동비} (원가 버전 C01)
    costs: dict[Key, UnitCost]
    #: {(plant_id, period): 월 고정비}
    fixed_costs: dict[tuple[str, date], Decimal]
    #: R01-v1 임계값 — `감소율_관심` 같은 키
    rules: dict[str, Decimal]
    models: list[str] = field(default_factory=list)


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _rows(ws: Any, first_key: str) -> Iterator[dict[str, Any]]:
    """`first_key`로 시작하는 헤더 행을 찾고, 그 아래 행들을 열 이름 → 값으로 돌려준다."""
    header: list[str] | None = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == first_key:
                header = [str(c) if c is not None else "" for c in row]
            continue
        if not any(v is not None for v in row):
            continue
        if row[0] == "rule_key":  # 버전기준 시트 안의 두 번째 표
            return
        yield dict(zip(header, row, strict=False))
    if header is None:
        raise ValueError(f"{ws.title}: '{first_key}' 헤더 행이 없다")


def _rules(ws: Any) -> dict[str, Decimal]:
    out: dict[str, Decimal] = {}
    inside = False
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "rule_key":
            inside = True
            continue
        if inside and row and row[0] is not None:
            out[str(row[0])] = Decimal(str(row[1]))
    return out


def _key(r: dict[str, Any]) -> Key:
    return (r["customer_id"], r["plant_id"], r["model_id"], _as_date(r["period"]))


@cache
def load_seed(path: Path = SEED_PATH) -> Seed:
    wb = load_workbook(path, data_only=True, read_only=True)

    versions = {r["version_id"]: r for r in _rows(wb["버전기준"], "version_id")}
    rules = _rules(wb["버전기준"])

    plans: dict[str, dict[Key, int]] = {}
    for r in _rows(wb["물량스냅샷"], "version_id"):
        plans.setdefault(r["version_id"], {})[_key(r)] = int(r["quantity"])

    customer_plans: dict[str, list[SourceRow]] = {}
    for r in _rows(wb["고객원천"], "source_id"):
        q = r["quantity"]
        customer_plans.setdefault(r["customer_plan_id"], []).append(
            SourceRow(
                source_id=r["source_id"],
                customer_plan_id=r["customer_plan_id"],
                customer_id=r["customer_id"],
                plant_id=r["plant_id"],
                model_id=r["model_id"],
                period=_as_date(r["period"]),
                quantity=None if q is None else int(q),
                unit=r["unit"],
                quantity_scope=r["quantity_scope"],
                allocation_scope=r["allocation_scope"],
            )
        )

    actuals: dict[Key, int] = {}
    cutoffs: dict[Key, date] = {}
    for r in _rows(wb["실적"], "actual_version"):
        actuals[_key(r)] = int(r["quantity"])
        cutoffs[_key(r)] = _as_date(r["cutoff_date"])

    costs = {
        _key(r): UnitCost(
            price=Decimal(str(r["unit_price"])),
            variable_cost=Decimal(str(r["unit_variable_cost"])),
            cm=Decimal(str(r["unit_cm"])),
        )
        for r in _rows(wb["단위원가"], "cost_version")
    }

    fixed = {
        (r["plant_id"], _as_date(r["period"])): Decimal(str(r["total_fixed"]))
        for r in _rows(wb["고정비"], "plant_id")
    }

    models = sorted({k[2] for k in plans["T01"]})
    wb.close()
    return Seed(
        versions=versions,
        plans=plans,
        customer_plans=customer_plans,
        actuals=actuals,
        actual_cutoffs=cutoffs,
        costs=costs,
        fixed_costs=fixed,
        rules=rules,
        models=models,
    )
